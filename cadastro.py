"""Sistema simples de cadastro de pessoas -> planilha .xlsx

Uso:  python3 cadastro.py
Dependência:  pip install openpyxl
"""
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Salva em ~/Documentos (Linux Mint em português); se não existir, usa a home
PASTA = Path.home() / "Documentos"
if not PASTA.exists():
    PASTA = Path.home()
ARQUIVO = PASTA / "cadastro_pessoas.xlsx"

CABECALHO = ["Nome", "Data de nascimento", "E-mail", "Telefone", "Cidade", "Cadastrado em"]


# ---------- Validação ----------
def pedir_texto(rotulo, obrigatorio=True):
    while True:
        valor = input(f"{rotulo}: ").strip()
        if valor or not obrigatorio:
            return valor
        print("  ! Campo obrigatório.")


def pedir_data(rotulo):
    while True:
        valor = input(f"{rotulo} (dd/mm/aaaa): ").strip()
        digitos = re.sub(r"\D", "", valor)   # tira barras, pontos, espaços...

        if len(digitos) != 8:
            print("  ! Digite 8 números, ex: 28081998.")
            continue

        try:
            data = datetime.strptime(digitos, "%d%m%Y")
        except ValueError:
            print("  ! Data inválida.")
            continue

        if data > datetime.now():
            print("  ! A data não pode estar no futuro.")
            continue

        return data.strftime("%d/%m/%Y")


def pedir_email(rotulo):
    while True:
        valor = input(f"{rotulo}: ").strip().lower()
        if re.fullmatch(r"[\w.+-]+@[\w-]+\.[\w.-]+", valor):
            return valor
        print("  ! E-mail inválido.")


def pedir_telefone(rotulo):
    while True:
        valor = input(f"{rotulo}: ").strip()
        digitos = re.sub(r"\D", "", valor)
        if 8 <= len(digitos) <= 15:
            return valor
        print("  ! Telefone inválido (use de 8 a 15 dígitos).")


# ---------- Planilha ----------
def abrir_planilha():
    """Abre a planilha existente ou cria uma nova com cabeçalho."""
    if ARQUIVO.exists():
        wb = load_workbook(ARQUIVO)
        return wb, wb.active

    wb = Workbook()
    ws = wb.active
    ws.title = "Cadastro"
    ws.append(CABECALHO)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="305496")
    ws.freeze_panes = "A2"
    return wb, ws


def email_ja_existe(ws, email):
    return any(
        linha[2] and str(linha[2]).lower() == email
        for linha in ws.iter_rows(min_row=2, values_only=True)
    )


def ajustar_colunas(ws):
    for col in ws.columns:
        largura = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = largura + 3


def salvar(wb):
    try:
        wb.save(ARQUIVO)
    except PermissionError:
        print(f"\n! Não consegui salvar. Feche o arquivo {ARQUIVO} se estiver aberto no LibreOffice.")
        sys.exit(1)


# ---------- Programa ----------
def main():
    wb, ws = abrir_planilha()
    print(f"Arquivo: {ARQUIVO}\n(Ctrl+C a qualquer momento para sair)\n")

    try:
        while True:
            print("--- Novo cadastro ---")
            nome = pedir_texto("Nome completo").title()
            nascimento = pedir_data("Data de nascimento")
            email = pedir_email("E-mail")

            if email_ja_existe(ws, email):
                print("  ! Já existe alguém cadastrado com esse e-mail.\n")
                continue

            telefone = pedir_telefone("Telefone")
            cidade = pedir_texto("Cidade", obrigatorio=False).title()

            ws.append([nome, nascimento, email, telefone, cidade,
                       datetime.now().strftime("%d/%m/%Y %H:%M")])
            ajustar_colunas(ws)
            salvar(wb)  # salva a cada cadastro: nada se perde se o programa fechar
            print(f"✔ {nome} cadastrado(a). Total: {ws.max_row - 1}\n")

            if input("Cadastrar outra pessoa? [S/n]: ").strip().lower() == "n":
                break
    except (KeyboardInterrupt, EOFError):
        print("\nEncerrado.")

    print(f"Dados salvos em {ARQUIVO}")


if __name__ == "__main__":
    main()